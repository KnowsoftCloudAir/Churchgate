"""PPTX builder — text, images and native charts."""
from __future__ import annotations
import io
import json
import zipfile
from pathlib import Path
from typing import Any, List, Optional
from xml.sax.saxutils import escape


def _parse_chart_data(raw: str):
    labels, values = [], []
    if not raw:
        return labels, values
    try:
        data = json.loads(raw)
        if isinstance(data, dict) and "labels" in data:
            return list(data.get("labels") or []), [float(x) for x in (data.get("values") or [])]
    except Exception:
        pass
    for part in str(raw).split(","):
        if ":" in part:
            a, b = part.rsplit(":", 1)
            try:
                values.append(float(b.strip()))
                labels.append(a.strip()[:40])
            except ValueError:
                continue
    return labels, values


def _resolve_image_path(path: Optional[str], base: Optional[Path] = None) -> Optional[Path]:
    if not path:
        return None
    p = str(path).strip()
    if p.startswith("http://") or p.startswith("https://"):
        return None
    if p.startswith("/static/"):
        rel = p[len("/static/"):]
        candidates = []
        if base:
            candidates.append(base / "app" / "static" / rel)
            candidates.append(base / "static" / rel)
        candidates.append(Path("app/static") / rel)
        for c in candidates:
            if c.exists():
                return c
        return None
    fp = Path(p)
    return fp if fp.exists() else None


def build_pptx_bytes(title: str, slides: List[Any], base_dir: Optional[Path] = None) -> bytes:
    try:
        return _via_python_pptx(slides, base_dir=base_dir)
    except Exception as e:
        print("python-pptx path failed:", e)
        return _via_simple(title, slides)


def _via_python_pptx(slides: List[Any], base_dir: Optional[Path] = None) -> bytes:
    from pptx import Presentation as PP
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor as RgbColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.chart.data import CategoryChartData

    prs = PP()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    base = base_dir or Path(__file__).resolve().parent.parent

    for s in slides:
        slide = prs.slides.add_slide(blank)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        shape.fill.solid()
        bg = (getattr(s, "bg_color", None) or "#0f172a").lstrip("#")
        try:
            shape.fill.fore_color.rgb = RgbColor.from_string(bg[:6])
        except Exception:
            shape.fill.fore_color.rgb = RgbColor(0x0F, 0x17, 0x2A)
        try:
            shape.line.fill.background()
        except Exception:
            pass

        accent = (getattr(s, "accent", None) or "#14b8a6").lstrip("#")
        try:
            accent_rgb = RgbColor.from_string(accent[:6])
        except Exception:
            accent_rgb = RgbColor(0x5E, 0xEA, 0xD4)

        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.2), Inches(1.0))
        run = title_box.text_frame.paragraphs[0].add_run()
        run.text = (getattr(s, "title", None) or "Slide")[:200]
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = accent_rgb

        body_text = getattr(s, "body", None) or ""
        body_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(7.0), Inches(4.5))
        bf = body_box.text_frame
        bf.word_wrap = True
        lines = body_text.splitlines() or [""]
        for i, line in enumerate(lines[:18]):
            para = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
            run = para.add_run()
            run.text = line[:400]
            run.font.size = Pt(16)
            run.font.color.rgb = RgbColor(0xE2, 0xE8, 0xF0)

        img_path = getattr(s, "image_path", None)
        images_json = getattr(s, "images_json", None)
        paths = []
        if images_json:
            try:
                arr = json.loads(images_json) if isinstance(images_json, str) else images_json
                if isinstance(arr, list):
                    paths.extend(arr)
            except Exception:
                pass
        if img_path:
            paths.insert(0, img_path)
        placed = False
        for up in paths[:1]:
            fp = _resolve_image_path(up, base)
            if fp:
                try:
                    slide.shapes.add_picture(str(fp), Inches(8.2), Inches(1.5), width=Inches(4.5))
                    placed = True
                except Exception as e:
                    print("img embed fail", e)

        chart_type = (getattr(s, "chart_type", None) or "").lower()
        labels, values = _parse_chart_data(getattr(s, "chart_data", None) or "")
        if chart_type and labels and values:
            try:
                cd = CategoryChartData()
                cd.categories = labels[:12]
                cd.add_series(getattr(s, "title", None) or "Series", values[:12])
                type_map = {
                    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
                    "line": XL_CHART_TYPE.LINE_MARKERS,
                    "pie": XL_CHART_TYPE.PIE,
                    "doughnut": XL_CHART_TYPE.DOUGHNUT,
                    "area": XL_CHART_TYPE.AREA,
                }
                xl_type = type_map.get(chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)
                if chart_type in ("pie", "doughnut"):
                    left, top, width, height = Inches(8.0), Inches(2.8), Inches(4.5), Inches(3.5)
                else:
                    left = Inches(7.8 if placed else 7.5)
                    top = Inches(4.0 if placed else 3.0)
                    width, height = Inches(5.0), Inches(3.0)
                slide.shapes.add_chart(xl_type, left, top, width, height, cd)
            except Exception as e:
                print("chart embed fail", e)

        foot = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12), Inches(0.35))
        fr = foot.text_frame.paragraphs[0].add_run()
        fr.text = "Knowsoft Eleon"
        fr.font.size = Pt(10)
        fr.font.color.rgb = RgbColor(0x64, 0x74, 0x8B)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _via_simple(deck_title: str, slides: List[Any]) -> bytes:
    """Minimal text-only fallback."""
    from pptx import Presentation as PP
    from pptx.util import Inches, Pt
    prs = PP()
    blank = prs.slide_layouts[6]
    for s in slides:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = (getattr(s, "title", None) or "Slide")[:200]
        run.font.size = Pt(28)
        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(5))
        bf = body.text_frame
        bf.word_wrap = True
        for i, line in enumerate((getattr(s, "body", None) or "").splitlines()[:20] or [""]):
            para = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
            para.add_run().text = line[:400]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
