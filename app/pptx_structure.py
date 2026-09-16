"""Extract slide structure from PPTX without destroying the original file."""
from __future__ import annotations
import io
import json
from typing import Any, List


def extract_pptx_structure(data: bytes, max_slides: int = 50) -> List[dict]:
    """Return list of {index, title, body, eleon_speak} preserving order."""
    out: List[dict] = []
    try:
        from pptx import Presentation as PP
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = PP(io.BytesIO(data))
        for i, slide in enumerate(prs.slides):
            if i >= max_slides:
                break
            texts = []
            title = ""
            for shape in slide.shapes:
                try:
                    if not shape.has_text_frame:
                        continue
                    t = (shape.text_frame.text or "").strip()
                    if not t:
                        continue
                    # first non-empty often title
                    if not title and len(t) < 120:
                        title = t.split("\n")[0][:120]
                    texts.append(t)
                except Exception:
                    continue
            # notes
            notes = ""
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            except Exception:
                pass
            body = "\n".join(texts[1:] if len(texts) > 1 else texts)[:4000]
            if not title:
                title = f"Slide {i + 1}"
            # default Eleon speak: notes or title + first lines
            speak = notes or (title + (". " + body[:400] if body else ""))
            out.append({
                "index": i,
                "title": title,
                "body": body,
                "notes": notes,
                "eleon_speak": speak[:2000],
            })
    except Exception as e:
        out.append({
            "index": 0,
            "title": "Import note",
            "body": f"Could not parse structure: {e}",
            "notes": "",
            "eleon_speak": "Welcome to this presentation.",
        })
    if not out:
        out.append({
            "index": 0,
            "title": "Slide 1",
            "body": "",
            "notes": "",
            "eleon_speak": "Welcome.",
        })
    return out


def scripts_to_json(scripts: List[dict]) -> str:
    return json.dumps(scripts, ensure_ascii=False)


def scripts_from_json(raw: Any) -> List[dict]:
    if not raw:
        return []
    if isinstance(raw, list):
        return raw
    try:
        data = json.loads(raw)
        return data if isinstance(data, list) else []
    except Exception:
        return []
